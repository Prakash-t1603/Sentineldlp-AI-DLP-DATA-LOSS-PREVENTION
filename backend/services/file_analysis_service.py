import io
import re
import csv
import json
import time
import zipfile
import tempfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Optional, Tuple, Dict, Any, List

from backend.utils.helpers import get_logger

logger = get_logger("SentinelDLP.FileAnalysisService")

# Supported Document Extensions for direct text extraction
DOCUMENT_EXTENSIONS = {
    ".txt", ".csv", ".json", ".md", ".log", ".py", ".js",
    ".html", ".css", ".xml", ".yaml", ".yml", ".env", ".sql",
    ".conf", ".ini", ".key", ".pem", ".cert", ".tsv", ".cfg",
    ".sh", ".bash", ".c", ".cpp", ".java", ".go", ".rs", ".php"
}
PDF_EXTENSIONS = {".pdf"}
DOCX_EXTENSIONS = {".docx", ".doc"}
PPTX_EXTENSIONS = {".pptx", ".ppt"}
XLSX_EXTENSIONS = {".xlsx", ".xls", ".xlsm"}
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".bmp", ".tiff", ".webp"}

# Max file size to analyze in memory (50MB)
MAX_ANALYSIS_FILE_SIZE = 50 * 1024 * 1024


class FileAnalysisService:
    """
    Centralized file content extraction service supporting multi-format text parsing,
    OCR for image documents and scanned PDFs, structured spreadsheet and CSV inspection,
    and large-file safety boundaries.
    """

    @staticmethod
    def is_supported_document(filepath: Path) -> bool:
        ext = filepath.suffix.lower()
        return ext in (DOCUMENT_EXTENSIONS | PDF_EXTENSIONS | DOCX_EXTENSIONS | PPTX_EXTENSIONS | XLSX_EXTENSIONS | IMAGE_EXTENSIONS)

    @staticmethod
    def extract_text_from_bytes(content_bytes: bytes, filename: str = "payload.dat") -> Tuple[str, str, Dict[str, Any]]:
        """
        Extract text from raw in-memory bytes based on filename extension.
        Returns (extracted_text, method_used, metadata_dict).
        """
        if not content_bytes:
            return "", "EMPTY_PAYLOAD", {}

        file_size = len(content_bytes)
        if file_size > MAX_ANALYSIS_FILE_SIZE:
            logger.warning(f"File {filename} exceeds max size limit ({file_size} bytes).")
            return "", "CONTENT_ANALYSIS_LIMIT_EXCEEDED", {"size": file_size, "limit": MAX_ANALYSIS_FILE_SIZE}

        ext = Path(filename).suffix.lower() if filename else ".dat"
        meta: Dict[str, Any] = {"file_size": file_size, "file_name": filename, "extension": ext}

        # 1. Plain Text, CSV, JSON, MD, Configs, Code
        if ext in DOCUMENT_EXTENSIONS:
            text = FileAnalysisService._extract_plain_text_from_bytes(content_bytes, ext)
            return text, "PLAIN_TEXT", meta

        # 2. PDF Documents (PyPDF text + embedded OCR fallback)
        elif ext in PDF_EXTENSIONS:
            from backend.ai.ocr_engine import ocr_engine
            ocr_res = ocr_engine.extract_from_pdf_bytes(content_bytes)
            meta["pages_processed"] = ocr_res.pages_processed
            meta["ocr_used"] = ocr_res.tier_used != "PYPDF_STREAM"
            meta["ocr_tier"] = ocr_res.tier_used
            return ocr_res.text or "", ocr_res.tier_used, meta

        # 3. Image OCR
        elif ext in IMAGE_EXTENSIONS:
            from backend.ai.ocr_engine import ocr_engine
            ocr_res = ocr_engine.extract_from_image_bytes(content_bytes, filename=filename)
            meta["ocr_used"] = True
            meta["ocr_confidence"] = ocr_res.confidence
            meta["ocr_tier"] = ocr_res.tier_used
            return ocr_res.text or "", "OCR", meta

        # 4. DOCX Documents (using temporary file for python-docx)
        elif ext in DOCX_EXTENSIONS:
            with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
                tmp.write(content_bytes)
                tmp_path = Path(tmp.name)
            try:
                text = FileAnalysisService._extract_docx(tmp_path)
                return text, "DOCX", meta
            finally:
                try:
                    tmp_path.unlink()
                except Exception:
                    pass

        # 5. XLSX Spreadsheets
        elif ext in XLSX_EXTENSIONS:
            with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
                tmp.write(content_bytes)
                tmp_path = Path(tmp.name)
            try:
                text = FileAnalysisService._extract_xlsx(tmp_path)
                return text, "OPENPYXL", meta
            finally:
                try:
                    tmp_path.unlink()
                except Exception:
                    pass

        # 6. PPTX Presentations
        elif ext in PPTX_EXTENSIONS:
            with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
                tmp.write(content_bytes)
                tmp_path = Path(tmp.name)
            try:
                text = FileAnalysisService._extract_pptx(tmp_path)
                return text, "PPTX", meta
            finally:
                try:
                    tmp_path.unlink()
                except Exception:
                    pass

        # Fallback: Attempt UTF-8 decode
        try:
            text = content_bytes.decode("utf-8", errors="ignore")[:100_000]
            return text, "RAW_DECODE", meta
        except Exception:
            return "", "UNSUPPORTED_FORMAT", meta

    @staticmethod
    def extract_text_from_file(filepath: Path, max_retries: int = 3) -> Tuple[str, str]:
        """
        Extract text from supported document types with retry handling for file write locks.
        Returns (extracted_text, method_used).
        """
        if not filepath or not filepath.exists() or not filepath.is_file():
            return "", "NOT_FOUND"

        for attempt in range(max_retries):
            try:
                file_size = filepath.stat().st_size
                if file_size == 0 and attempt < max_retries - 1:
                    time.sleep(0.3)
                    continue

                if file_size > MAX_ANALYSIS_FILE_SIZE:
                    logger.warning(f"File {filepath.name} exceeds max size ({file_size} bytes). Skipping deep extraction.")
                    return "", "CONTENT_ANALYSIS_LIMIT_EXCEEDED"

                ext = filepath.suffix.lower()

                # 1. Plain Text, CSV, JSON, MD, Configs, SQL, .env, Code
                if ext in DOCUMENT_EXTENSIONS:
                    text = FileAnalysisService._extract_plain_text(filepath)
                    if text or attempt == max_retries - 1:
                        return text, "PLAIN_TEXT"

                # 2. PDF Documents (with OCR fallback for scanned PDFs)
                elif ext in PDF_EXTENSIONS:
                    text = FileAnalysisService._extract_pdf(filepath)
                    if not text.strip():
                        from backend.ai.ocr_engine import ocr_engine
                        try:
                            with open(filepath, "rb") as f:
                                pdf_bytes = f.read()
                            ocr_res = ocr_engine.extract_from_pdf_bytes(pdf_bytes)
                            if ocr_res.text:
                                return ocr_res.text, "PDF_OCR"
                        except Exception:
                            pass
                    if text or attempt == max_retries - 1:
                        return text, "PYPDF"

                # 3. DOCX Documents
                elif ext in DOCX_EXTENSIONS:
                    text = FileAnalysisService._extract_docx(filepath)
                    if text or attempt == max_retries - 1:
                        return text, "DOCX"

                # 4. PPTX Presentations
                elif ext in PPTX_EXTENSIONS:
                    text = FileAnalysisService._extract_pptx(filepath)
                    if text or attempt == max_retries - 1:
                        return text, "PPTX"

                # 5. XLSX Spreadsheets
                elif ext in XLSX_EXTENSIONS:
                    text = FileAnalysisService._extract_xlsx(filepath)
                    if text or attempt == max_retries - 1:
                        return text, "OPENPYXL"

                # 6. Image documents handled by OCR
                elif ext in IMAGE_EXTENSIONS:
                    from backend.ai.ocr_engine import ocr_engine
                    try:
                        with open(filepath, "rb") as f:
                            img_bytes = f.read()
                        ocr_res = ocr_engine.extract_from_image_bytes(img_bytes, filename=filepath.name)
                        return ocr_res.text or "", "OCR"
                    except Exception:
                        from backend.services.ocr_service import ocr_service
                        text = ocr_service.extract_text_from_image(filepath)
                        return text, "OCR"

            except (PermissionError, OSError) as e:
                if attempt < max_retries - 1:
                    time.sleep(0.3 * (attempt + 1))
                    continue
                logger.debug(f"File extraction retry exhausted for {filepath.name}: {e}")

            except Exception as e:
                logger.warning(f"Unexpected extraction error for {filepath.name}: {e}")
                break

        return "", "UNSUPPORTED_OR_EMPTY"

    @staticmethod
    def _extract_plain_text(filepath: Path) -> str:
        ext = filepath.suffix.lower()
        for encoding in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
            try:
                with open(filepath, "r", encoding=encoding, errors="ignore") as f:
                    raw_content = f.read(250_000)
                    return FileAnalysisService._format_structured_text(raw_content, ext)
            except (PermissionError, OSError):
                raise
            except Exception:
                continue
        return ""

    @staticmethod
    def _extract_plain_text_from_bytes(content_bytes: bytes, ext: str) -> str:
        for encoding in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
            try:
                raw_content = content_bytes.decode(encoding, errors="ignore")[:250_000]
                return FileAnalysisService._format_structured_text(raw_content, ext)
            except Exception:
                continue
        return ""

    @staticmethod
    def _format_structured_text(raw_content: str, ext: str) -> str:
        """Format CSV, JSON, and text content for optimal NLP and entity recognition."""
        if not raw_content:
            return ""

        # CSV / TSV handling
        if ext in (".csv", ".tsv"):
            try:
                delimiter = "\t" if ext == ".tsv" or "\t" in raw_content[:500] else ","
                reader = csv.reader(io.StringIO(raw_content), delimiter=delimiter)
                lines = []
                headers = []
                for i, row in enumerate(reader):
                    if i == 0:
                        headers = [h.strip() for h in row]
                        lines.append(f"Headers: {' | '.join(headers)}")
                    elif i < 150:  # Process first 150 rows
                        row_items = []
                        for col_idx, val in enumerate(row):
                            val_clean = val.strip()
                            if val_clean:
                                col_name = headers[col_idx] if col_idx < len(headers) else f"Col_{col_idx+1}"
                                row_items.append(f"{col_name}: {val_clean}")
                        if row_items:
                            lines.append(" | ".join(row_items))
                if lines:
                    return "\n".join(lines)
            except Exception:
                pass

        # JSON handling
        elif ext == ".json":
            try:
                parsed = json.loads(raw_content)
                flattened = []
                def _walk(obj, prefix=""):
                    if isinstance(obj, dict):
                        for k, v in obj.items():
                            _walk(v, f"{prefix}.{k}" if prefix else str(k))
                    elif isinstance(obj, list):
                        for item in obj[:100]:
                            _walk(item, prefix)
                    else:
                        flattened.append(f"{prefix}: {obj}")
                _walk(parsed)
                if flattened:
                    return "\n".join(flattened[:300])
            except Exception:
                pass

        return raw_content

    @staticmethod
    def _extract_pdf(filepath: Path) -> str:
        try:
            from pypdf import PdfReader
            reader = PdfReader(str(filepath))
            text_parts = []
            for page in reader.pages[:30]:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
            return "\n".join(text_parts)
        except Exception as e:
            logger.debug(f"PDF extraction error for {filepath.name}: {e}")
            return ""

    @staticmethod
    def _extract_docx(filepath: Path) -> str:
        try:
            import docx
            doc = docx.Document(str(filepath))
            text_parts = []

            # Paragraphs
            for p in doc.paragraphs:
                if p.text.strip():
                    text_parts.append(p.text.strip())

            # Tables
            for table in doc.tables[:20]:
                for row in table.rows:
                    row_cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if row_cells:
                        text_parts.append(" | ".join(row_cells))

            # Headers & Footers
            for section in doc.sections:
                if section.header:
                    for hp in section.header.paragraphs:
                        if hp.text.strip():
                            text_parts.append(f"Header: {hp.text.strip()}")
                if section.footer:
                    for fp in section.footer.paragraphs:
                        if fp.text.strip():
                            text_parts.append(f"Footer: {fp.text.strip()}")

            return "\n".join(text_parts)
        except Exception as e:
            logger.debug(f"DOCX extraction error for {filepath.name}: {e}")
            return ""

    @staticmethod
    def _extract_pptx(filepath: Path) -> str:
        try:
            text_parts = []
            with zipfile.ZipFile(filepath, 'r') as z:
                slide_files = [f for f in z.namelist() if f.startswith("ppt/slides/slide") and f.endswith(".xml")]
                for slide in sorted(slide_files):
                    xml_content = z.read(slide)
                    tree = ET.fromstring(xml_content)
                    for elem in tree.iter():
                        if elem.tag.endswith('}t') and elem.text:
                            text_parts.append(elem.text.strip())
            return "\n".join(text_parts)
        except Exception as e:
            logger.debug(f"PPTX extraction error for {filepath.name}: {e}")
            return ""

    @staticmethod
    def _extract_xlsx(filepath: Path) -> str:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(str(filepath), data_only=True, read_only=True)
            text_parts = []
            for sheet_name in wb.sheetnames[:5]:  # Process first 5 sheets
                sheet = wb[sheet_name]
                headers = []
                for row_idx, row in enumerate(sheet.iter_rows(max_row=150, values_only=True)):
                    row_vals = [str(val).strip() for val in row if val is not None and str(val).strip()]
                    if not row_vals:
                        continue
                    if row_idx == 0:
                        headers = row_vals
                        text_parts.append(f"Sheet '{sheet_name}' Headers: {' | '.join(headers)}")
                    else:
                        row_items = []
                        for col_idx, val in enumerate(row):
                            if val is not None and str(val).strip():
                                col_title = headers[col_idx] if col_idx < len(headers) else f"Col_{col_idx+1}"
                                row_items.append(f"{col_title}: {str(val).strip()}")
                        if row_items:
                            text_parts.append(" | ".join(row_items))
                        else:
                            text_parts.append(" | ".join(row_vals))
            return "\n".join(text_parts)
        except Exception as e:
            logger.debug(f"XLSX extraction error for {filepath.name}: {e}")
            return ""

    @staticmethod
    def analyze_with_ai(
        db,
        filepath: Path,
        employee_id: Optional[str] = None,
        channel: str = "FILE",
        destination: str = "",
        device_id: str = "WORKSTATION",
        event_id: Optional[str] = None
    ):
        """
        Analyze document or image using Central AI Pipeline.
        """
        from backend.ai.pipeline import dlp_pipeline
        extracted_text, method = FileAnalysisService.extract_text_from_file(filepath)

        content_bytes = None
        try:
            if filepath.exists() and filepath.stat().st_size <= MAX_ANALYSIS_FILE_SIZE:
                with open(filepath, "rb") as f:
                    content_bytes = f.read()
        except Exception:
            pass

        return dlp_pipeline.analyze_payload(
            db=db,
            content=extracted_text,
            content_bytes=content_bytes,
            filename=filepath.name,
            channel=channel,
            destination=destination,
            employee_id=employee_id,
            device_id=device_id,
            event_id=event_id,
            persist=True
        )


file_analysis_service = FileAnalysisService()
